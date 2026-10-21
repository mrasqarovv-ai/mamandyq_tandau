#!/usr/bin/env python3
"""Export retrieved Daryn documents, with their full text and provenance, to Word.

Run from the repository root:
    python3 tools/export_docx.py --archive data/archive.json \
        --output downloads/daryn-archive.docx

This command deliberately fails for an empty or malformed archive. It never
creates placeholder questions or represents unchecked years as complete.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from urllib.parse import urlsplit


SUBJECTS = {
    "mathematics": "Математика",
    "physics": "Физика",
    "chemistry": "Химия",
    "biology": "Биология",
    "geography": "География",
    "informatics": "Информатика",
    "history": "Қазақстан тарихы",
    "kazakh": "Қазақ тілі мен әдебиеті",
    "russian": "Орыс тілі мен әдебиеті",
    "english": "Ағылшын тілі",
    "natural-science": "Жаратылыстану",
}
LANGUAGES = {"kk": "Қазақша", "ru": "Орысша", "mixed": "Аралас тілдер",
             "kk-ru": "Қазақша / орысша", "en": "Ағылшынша",
             "unknown": "Бастапқы тілін тексеру қажет"}
COVERAGE = {
    "unverified": "Тексерілмеген",
    "partial": "Ішінара жиналған",
    "checked": "Қолжетімді материалдары тексерілген",
}
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def string(value: object, field: str, *, required: bool = True) -> str:
    if not isinstance(value, str) or (required and not value.strip()):
        raise ValueError(f"{field}: бос емес мәтін болуы керек")
    if INVALID_XML.search(value):
        raise ValueError(f"{field}: Word құжатына жарамсыз басқару таңбалары бар")
    return value


def source_url(value: object, field: str, *, official: bool = False) -> str:
    result = string(value, field)
    parsed = urlsplit(result)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError(f"{field}: толық HTTP(S) дереккөз сілтемесі қажет")
    if parsed.username is not None or parsed.password is not None:
        raise ValueError(f"{field}: сілтемеде құпия деректер болмауы керек")
    if official and parsed.hostname.lower() not in {"daryn.kz", "www.daryn.kz"}:
        raise ValueError(f"{field}: ресми daryn.kz бетіне сілтеме қажет")
    return result


def validate_archive(raw: object) -> dict:
    if not isinstance(raw, dict) or raw.get("schemaVersion") != 1:
        raise ValueError("schemaVersion: қолдау көрсетілетін нұсқа — 1")
    if raw.get("status") not in {"source-access-blocked", "partial", "ready"}:
        raise ValueError("Архивтің status өрісі жарамсыз")
    string(raw.get("updatedAt"), "updatedAt")
    documents = raw.get("documents")
    if not isinstance(documents, list) or not documents:
        raise ValueError(
            "Нақты тапсырмалар әлі жүктелмеген: documents бос. "
            "Бос немесе ойдан шығарылған Word архиві жасалмады."
        )
    identifiers: set[str] = set()
    for index, record in enumerate(documents):
        field = f"documents[{index}]"
        if not isinstance(record, dict):
            raise ValueError(f"{field}: объект болуы керек")
        for key in ("id", "title", "subject", "year", "stage"):
            string(record.get(key), f"{field}.{key}")
        string(record.get("text", ""), f"{field}.text", required=False)
        if record["id"] in identifiers:
            raise ValueError(f"{field}.id: қайталанған құжат идентификаторы")
        identifiers.add(record["id"])
        if record.get("grade") is not None and (
            type(record["grade"]) is not int or record["grade"] not in range(5, 9)
        ):
            raise ValueError(f"{field}.grade: 5–8 аралығындағы сынып немесе null қажет")
        grades = record.get("grades", [])
        if not isinstance(grades, list) or any(type(g) is not int or g not in range(5, 9) for g in grades):
            raise ValueError(f"{field}.grades: 5–8 сыныптарының тізімі болуы керек")
        if record.get("language") not in LANGUAGES:
            raise ValueError(f"{field}.language: қолдау көрсетілмейтін тіл белгісі")
        source_url(record.get("sourcePage"), f"{field}.sourcePage", official=True)
        source_url(record.get("sourceUrl"), f"{field}.sourceUrl")
        questions = record.get("questions", [])
        if not isinstance(questions, list):
            raise ValueError(f"{field}.questions: тізім болуы керек")
        for qi, question in enumerate(questions):
            if not isinstance(question, dict):
                raise ValueError(f"{field}.questions[{qi}]: объект болуы керек")
            string(question.get("text"), f"{field}.questions[{qi}].text")
    coverage = raw.get("coverage", [])
    if not isinstance(coverage, list):
        raise ValueError("coverage: тізім болуы керек")
    for index, item in enumerate(coverage):
        if not isinstance(item, dict) or item.get("status") not in COVERAGE:
            raise ValueError(f"coverage[{index}]: мәртебесі жарамсыз")
        string(item.get("year"), f"coverage[{index}].year")
        string(item.get("note", ""), f"coverage[{index}].note", required=False)
    return raw


def add_link(paragraph, label: str, url: str, *, short: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    if not short:
        paragraph.add_run(label + ": ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set("{%s}id" % "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                  paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    properties.append(color)
    font = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        font.set(qn("w:" + attribute), "Times New Roman")
    properties.append(font)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    properties.append(size)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label if short else url
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def record_grades(record: dict) -> list[int]:
    if record.get("grade") is not None:
        return [record["grade"]]
    return sorted(set(record.get("grades", [])))


def grade_label(record: dict) -> str:
    grades = record_grades(record)
    return ", ".join(map(str, grades)) + "-сынып" if grades else "Сыныбы тексерілмеген"


def task_count(record: dict) -> str:
    questions = record.get("questions", [])
    if questions:
        suffix = "" if record.get("questionCountVerified") is True else " (бөлінген; толық саны тексерілмеген)"
        return f"{len(questions)}{suffix}"
    return "Нақты саны тексерілмеген"


def material_type(record: dict) -> str:
    return record.get("materialType", record.get("material_type", "tasks"))


def sort_key(record: dict) -> tuple:
    stage = record["stage"].casefold()
    stage_order = 0 if any(w in stage for w in ("аудан", "район", "қалалық", "город")) else (
        1 if any(w in stage for w in ("облыс", "област")) else 2 if "респ" in stage else 3)
    language_order = {"kk": 0, "kk-ru": 1, "mixed": 1, "ru": 2, "en": 3, "unknown": 4}
    return (record["year"], stage_order, stage, SUBJECTS.get(record["subject"], record["subject"]),
            language_order[record["language"]], record["id"])


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = " " + instruction + " "
    run._r.append(code)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    run._r.append(separator)
    if placeholder:
        paragraph.add_run(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def original_path(record: dict, archive_path: Path | None) -> Path | None:
    if archive_path is None or not isinstance(record.get("fileUrl"), str):
        return None
    relative = Path(record["fileUrl"])
    if relative.is_absolute() or ".." in relative.parts:
        return None
    for root in (archive_path.parent.parent.resolve(), archive_path.parent.resolve()):
        candidate = (root / relative).resolve()
        if candidate.is_relative_to(root) and candidate.is_file():
            return candidate
    return None


def add_original_text(document, record: dict, archive_path: Path | None) -> None:
    text = record.get("text", "")
    if text.strip():
        # Keep all source text. PDF form-feeds become genuine Word page breaks.
        pages = text.replace("\r\n", "\n").replace("\r", "\n").split("\f")
        for index, page in enumerate(pages):
            if index:
                document.add_page_break()
            for block in page.split("\n\n"):
                document.add_paragraph(block)
        return
    path = original_path(record, archive_path)
    if path is not None and path.suffix.lower() == ".pdf":
        import io
        import fitz
        from docx.shared import Cm

        document.add_paragraph("Мәтін қабаты алынбаған. Төменде түпнұсқа PDF беттерінің көшірмелері берілген.")
        with fitz.open(path) as original:
            for page in original:
                picture = page.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
                document.add_picture(io.BytesIO(picture.tobytes("png")), width=Cm(16.6))
        return
    document.add_paragraph(
        "Бұл құжаттың мәтіні алынбаған; осы Word жинағында тапсырмалардың толық мазмұны жоқ. "
        "Түпнұсқаны берілген сілтеме немесе сайттағы жүктеу батырмасы арқылы ашыңыз."
    )


def add_record(document, record: dict, archive_path: Path | None) -> None:
    document.add_heading(record["title"], 3)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    fields = (
        ("Оқу жылы", record["year"]), ("Сынып", grade_label(record)),
        ("Пән", SUBJECTS.get(record["subject"], record["subject"])),
        ("Олимпиада кезеңі", record["stage"]), ("Тапсырма саны", task_count(record)),
        ("Бастапқы тілі", LANGUAGES[record["language"]]),
        ("Материал түрі", {"answers": "Жарияланған жауап кілті", "solutions": "Жарияланған шешімдер"}.get(material_type(record), "Тапсырмалар")),
    )
    for label, value in fields:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    add_link(document.add_paragraph(), "Ресми бет", record["sourcePage"])
    add_link(document.add_paragraph(), "Түпнұсқа құжат", record["sourceUrl"])
    document.add_paragraph(f"Архив идентификаторы: {record['id']}")
    if record.get("extractionNote"):
        document.add_paragraph(str(record["extractionNote"]))
    document.add_heading("Түпнұсқадан алынған мәтін", 4)
    add_original_text(document, record, archive_path)


def add_record_group(document, records: list[dict], archive_path: Path | None) -> None:
    previous_year = None
    for record in sorted(records, key=sort_key):
        if record["year"] != previous_year:
            document.add_heading(record["year"] + " оқу жылы", 2)
            previous_year = record["year"]
        add_record(document, record, archive_path)


def build_document(archive: dict, output: Path, archive_path: Path | None = None) -> None:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for style in document.styles:
        if style.type in (1, 2):
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
    for style_name, size in (("Title", 18), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 12)):
        style = document.styles[style_name]
        style.font.size, style.font.bold = Pt(size), True
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(section.footer.paragraphs[0], "PAGE", "1")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    document.core_properties.title = "Daryn · Олимпиада тапсырмалары"
    document.core_properties.subject = "5–8-сыныптарға арналған ресми тапсырмалар жинағы"
    document.add_heading("5–8-СЫНЫПТАРҒА АРНАЛҒАН РЕСПУБЛИКАЛЫҚ ОЛИМПИАДА ТАПСЫРМАЛАРЫНЫҢ ЖИНАҒЫ (2020–2027)", 0)
    document.add_paragraph("Ресми дереккөздерден нақты алынған материалдар · Ішінара архив")
    document.add_paragraph(f"Архив жаңартылған уақыт: {archive['updatedAt']}")
    document.add_paragraph(f"Жинақтағы құжат жазбаларының саны: {len(archive['documents'])}")
    document.add_paragraph(
        "Бұл басылым қазақша, орысша, ағылшынша және аралас тілдегі түпнұсқаларды сақтайды. "
        "Мәтіндер аударылмаған; қазақша нұсқалар бар топтамаларда алдымен көрсетіледі."
    )
    document.add_paragraph(
        "Бұл жинаққа тек архивке нақты жүктелген құжаттардың мәтіні енгізілді. "
        "Жылдың тексерілуі барлық тапсырмалардың толық жиналғанын білдірмейді. "
        "Бұл — түпнұсқа мәтіндерінің транскрипциясы: суреттер, диаграммалар, кестелердің "
        "ұяшықтары, математикалық дәрежелер мен индекстер автоматты мәтінге толық көшпеуі мүмкін. "
        "Формулалар мен графикалық есептерді міндетті түрде сайттағы PDF/DOCX түпнұсқасымен "
        "салыстырыңыз. Түпнұсқа файлдар бөлек сақталған; мәтінсіз PDF беттері қолжетімді болса "
        "ғана осы құжатқа сурет ретінде енгізіледі."
    )
    if archive.get("status") == "source-access-blocked":
        document.add_paragraph(
            "Дереккөздерге қолжетімділік шектелген; архивті толықтыру аяқталған жоқ."
        )
    document.add_page_break()
    document.add_heading("Мазмұны", 1)
    add_field(document.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
              "Автоматты мазмұнды Word бағдарламасында өрісті жаңарту арқылы көрсетіңіз (Ctrl+A, F9).")
    document.add_page_break()
    document.add_heading("Жинақтың қорытынды кестесі", 1)
    document.add_paragraph(
        "«Бөлінген» саны — сұрақтарға сенімді бөлініп сақталған жазбалар саны; "
        "қалған мәтіндегі тапсырмалардың толық саны емес. Нөлдік тізім тапсырма жоқ дегенді "
        "білдірмейді. Бүкіл файлдары тексерілмеген жылдар толық деп белгіленбейді."
    )
    summary = document.add_table(rows=1, cols=7)
    summary.style = "Table Grid"
    for cell, label in zip(summary.rows[0].cells, ("Оқу жылы", "Сынып", "Кезең", "Пән", "Тапсырма саны", "Дереккөз", "Мәртебесі")):
        cell.text = label
    ordered = sorted(archive["documents"], key=lambda r: (record_grades(r) or [9], sort_key(r)))
    for record in ordered:
        cells = summary.add_row().cells
        values = (record["year"], grade_label(record), record["stage"], SUBJECTS.get(record["subject"], record["subject"]), task_count(record))
        for cell, value in zip(cells, values):
            cell.text = value
        add_link(cells[5].paragraphs[0], "Түпнұсқа", record["sourceUrl"], short=True)
        cells[6].text = "Жартылай жинақталды" if record.get("text", "").strip() else "Тексеруді қажет етеді"
    tasks = [r for r in ordered if material_type(r) not in {"answers", "solutions"}]
    answers = [r for r in ordered if material_type(r) in {"answers", "solutions"}]
    roman = {5: "І", 6: "ІІ", 7: "ІІІ", 8: "IV"}
    for grade in range(5, 9):
        document.add_page_break()
        document.add_heading(f"{roman[grade]} бөлім. {grade}-сынып олимпиада тапсырмалары", 1)
        direct = [r for r in tasks if record_grades(r) == [grade]]
        if direct:
            add_record_group(document, direct, archive_path)
        else:
            document.add_paragraph("Бұл сыныпқа жеке расталған топтама енгізілмеген. Бұл тапсырмалар жарияланбаған деген қорытынды емес.")
        shared = [r for r in tasks if grade in record_grades(r) and len(record_grades(r)) > 1]
        if shared:
            document.add_heading("Бірнеше сыныпқа ортақ материалдарға сілтемелер", 2)
            document.add_paragraph("Мәтін қайталанбауы үшін бұл құжаттар ортақ материалдар қосымшасында бір рет берілген.")
            for record in shared:
                document.add_paragraph(f"{record['title']} · Архив ID: {record['id']}")
    shared = [r for r in tasks if len(record_grades(r)) > 1]
    unknown = [r for r in tasks if not record_grades(r)]
    for heading, group in (("Қосымша. Бірнеше сыныпқа ортақ тапсырмалар", shared),
                           ("Қосымша. Сыныбы нақты анықталмаған тапсырмалар", unknown)):
        if group:
            document.add_page_break()
            document.add_heading(heading, 1)
            add_record_group(document, group, archive_path)
    document.add_page_break()
    document.add_heading("V бөлім. Жарияланған ресми жауап кілттері", 1)
    document.add_paragraph(
        "Бұл бөлімге тек бастапқы архивте жауап/шешім деп белгіленген файлдар енгізілді. "
        "ЖИ жауаптары қосылмаған. Кілттің нақты нұсқаға сәйкестігін сыныбы, жылы, кезеңі "
        "және түпнұсқа атауы арқылы тексеріңіз; автоматты сәйкестендіру расталмаған."
    )
    if answers:
        add_record_group(document, answers, archive_path)
    else:
        document.add_paragraph("Расталған жауап кілттері бұл жинаққа әлі енгізілмеген.")
    document.add_page_break()
    document.add_heading("VI бөлім. Пайдаланылған дереккөздер және толықтық есебі", 1)
    document.add_heading("Оқу жылдарының қамтылуы", 2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ("Оқу жылы", "Мәртебе", "Ескерту")):
        cell.text = title
    coverage = {item["year"].replace("-", "–"): item for item in archive.get("coverage", [])}
    years = [f"{year}–{year + 1}" for year in range(2020, 2027)]
    years += sorted(set(coverage) - set(years))
    for year in years:
        item = coverage.get(year, {"status": "unverified", "note": "Бұл оқу жылының толықтығы тексерілмеген."})
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = year, COVERAGE[item["status"]], item.get("note", "")
    document.add_heading("Дереккөздер", 2)
    source_pages = sorted(set(record["sourcePage"] for record in ordered))
    for url in source_pages:
        add_link(document.add_paragraph(), "Ресми архив беті", url)
    for source in archive.get("sources", []):
        if isinstance(source, dict) and source.get("url") not in source_pages and source.get("url"):
            add_link(document.add_paragraph(), str(source.get("title", "Дереккөз")), source_url(source["url"], "sources.url"))
    document.add_paragraph(
        "Әр материалдың нақты бастапқы файлына сілтеме оның метадеректерінде берілді. "
        "Толық жүктеу журналы, қолжетімсіз файлдар мен бастапқы Drive бумаларының мәртебелері "
        "сайттағы дереккөздер есебінде сақталған. Бұл Word құжаты оларды тексерілді деп алмастырмайды."
    )
    document.add_heading("Тексеру шектері", 2)
    document.add_paragraph(
        "Архив ішінара: барлық жеті жыл, төрт сынып және үш кезең бойынша толықтық расталмаған. "
        "Желілік қате немесе бос іздеу нәтижесі материал жарияланбағанының дәлелі емес. "
        "Белгісіз сынып, тіл және толық сұрақ саны ойдан толтырылмаған. Бір PDF-тің әр тілдегі "
        "немесе әр сыныптағы үзінділері жеке жазба болуы мүмкін; сондықтан жазба саны түпнұсқа "
        "файлдар санымен бірдей болмауы мүмкін. Мәтін автоматты алынған, ресми емлесі мен "
        "формулалары қолмен түгел тексерілмеген; бастапқы мазмұн түзетілмеген."
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".daryn-", suffix=".docx", dir=output.parent)
    os.close(descriptor)
    try:
        document.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", type=Path, default=Path("data/archive.json"))
    parser.add_argument("--output", type=Path, default=Path("downloads/daryn-archive.docx"))
    args = parser.parse_args()
    try:
        with args.archive.open(encoding="utf-8") as handle:
            archive = validate_archive(json.load(handle))
        if args.output.suffix.lower() != ".docx":
            raise ValueError("Нәтиже файлының кеңейтімі .docx болуы керек")
        if args.archive.resolve() == args.output.resolve():
            raise ValueError("Архивтің үстіне жазуға болмайды")
        build_document(archive, args.output, args.archive)
    except (OSError, ValueError, ImportError) as error:
        print(f"Word экспорты орындалмады: {error}", file=sys.stderr)
        return 1
    print(f"Word жинағы сақталды: {args.output} ({len(archive['documents'])} құжат)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
